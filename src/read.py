from pathlib import Path

from openpyxl import load_workbook
from rich import print
from rich.markdown import Markdown

ROOT_FOLDER = Path(__file__).parent
WORKBOOKS_FOLDER = ROOT_FOLDER / "workbooks"
Path.mkdir(WORKBOOKS_FOLDER, exist_ok=True)

WORKBOOK_PATH = WORKBOOKS_FOLDER / "workbook.xlsx"

# loading a xlsx file
workbook = load_workbook(WORKBOOK_PATH)
sheet_name = "My Sheet"
worksheet = workbook[sheet_name]

# showing cell values
for row in worksheet.iter_rows(values_only=True):
    print(row)

print(Markdown("---"))

# show starting from the second row
for row in worksheet.iter_rows(min_row=2, values_only=True):
    print(row)

print(Markdown("---"))

# editting data from a cell
print(worksheet["b3"].value)
worksheet["b3"].value = 55
print(worksheet["b3"].value)

print(Markdown("---"))

# changing values inside a for loop
for row in worksheet.iter_rows(min_row=2):
    for cell in row:
        if cell.value == "Guina":
            # cell(row: int, col: int, value='Any value you want to change')
            worksheet.cell(cell.row, 2, 45)
            print(cell.value)

print(Markdown("---"))

workbook.save(WORKBOOK_PATH)
